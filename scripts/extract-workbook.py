"""Read-only XLSX extraction with source cells, cached values, formulas and formatting."""
import sys, json, hashlib
from pathlib import Path
from datetime import datetime, date
import openpyxl

def json_value(v):
    if isinstance(v, (datetime, date)): return v.isoformat()
    if v is None or isinstance(v, (str, int, float, bool)): return v
    return str(v)

def color(c):
    if c is None: return None
    return {"type": c.type, "value": str(c.value), "tint": c.tint}

def main():
    source=Path(sys.argv[1]).resolve()
    destination=Path(sys.argv[2] if len(sys.argv)>2 else ".local/workbook.json")
    formulas=openpyxl.load_workbook(source, data_only=False)
    cached=openpyxl.load_workbook(source, data_only=True)
    sheets=[]
    for sheet in formulas:
        rows=[]
        for row in sheet.iter_rows():
            cells=[]
            for cell in row:
                if cell.value is None: continue
                cells.append({"cell":cell.coordinate,"row":cell.row,"column":cell.column,
                    "value":json_value(cached[sheet.title][cell.coordinate].value if cell.data_type=="f" else cell.value),
                    "formula":(cell.value if isinstance(cell.value,str) else getattr(cell.value,"text",str(cell.value))) if cell.data_type=="f" else None,
                    "formula_range":getattr(cell.value,"ref",None) if cell.data_type=="f" else None,
                    "fill":color(cell.fill.fgColor) if cell.fill.patternType else None,
                    "font_color":color(cell.font.color),"bold":bool(cell.font.bold),"strike":bool(cell.font.strike),
                    "comment":cell.comment.text if cell.comment else None})
            if cells: rows.append({"row":row[0].row,"cells":cells})
        sheets.append({"name":sheet.title,"state":sheet.sheet_state,"rows":rows,"max_row":sheet.max_row,"max_column":sheet.max_column,
            "merged_ranges":[str(r) for r in sheet.merged_cells.ranges]})
    out={"source_filename":source.name,"sha256":hashlib.sha256(source.read_bytes()).hexdigest(),"sheets":sheets}
    destination.parent.mkdir(parents=True,exist_ok=True)
    destination.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding="utf-8")
    inventory=[{"name":s["name"],"rows":len(s["rows"]),"columns":s["max_column"],"formula_count":sum(bool(c["formula"]) for r in s["rows"] for c in r["cells"]),"sample":[[c["cell"],str(c["value"])[:150]] for r in s["rows"][:4] for c in r["cells"][:10]]} for s in sheets]
    destination.with_name("workbook-inventory.json").write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(inventory,ensure_ascii=False))
if __name__=="__main__": main()
